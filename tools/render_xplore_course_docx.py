from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ACCENT = RGBColor(26, 91, 147)
TEXT = RGBColor(36, 43, 51)
MUTED = RGBColor(89, 99, 110)
FONT_NAME = "Aptos"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, *, size: float = 11, bold: bool = False, color: RGBColor = TEXT, name: str = FONT_NAME) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, space_before, space_after in (
        ("Title", 22, 0, 14),
        ("Heading 1", 16, 18, 8),
        ("Heading 2", 13, 14, 6),
        ("Heading 3", 11.5, 10, 4),
    ):
        style = doc.styles[name]
        style.font.name = FONT_NAME
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(space_after)
        style.paragraph_format.keep_with_next = True

    code_style = doc.styles.add_style("Course Code", WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(9)
    code_style.font.color.rgb = TEXT
    code_style.paragraph_format.space_before = Pt(6)
    code_style.paragraph_format.space_after = Pt(6)
    code_style.paragraph_format.line_spacing = 1.0

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Softengi | Управління вимогами та змінами в SDLC")
    set_font(footer_run, size=8.5, color=MUTED)
    footer.add_run("  |  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def add_run_with_inline_bold(paragraph, text: str, *, font_size: float = 11, color: RGBColor = TEXT) -> None:
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        content = part[2:-2] if is_bold else part
        run = paragraph.add_run(content)
        set_font(run, size=font_size, bold=is_bold, color=color)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.58 + level * 0.48)
    paragraph.paragraph_format.first_line_indent = Cm(-0.42)
    bullet = paragraph.add_run("• ")
    set_font(bullet, size=11, bold=True, color=ACCENT)
    add_run_with_inline_bold(paragraph, text)


def add_numbered(doc: Document, label: str, text: str, level: int = 0) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.62 + level * 0.48)
    paragraph.paragraph_format.first_line_indent = Cm(-0.48)
    number = paragraph.add_run(f"{label} ")
    set_font(number, size=11, bold=True, color=ACCENT)
    add_run_with_inline_bold(paragraph, text)


def is_table_separator(line: str) -> bool:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        line = lines[index].strip()
        if not is_table_separator(line):
            rows.append([cell.strip() for cell in line.strip("|").split("|")])
        index += 1
    return rows, index


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=columns)
    table.style = "Table Grid"
    table.autofit = True

    for row_index, row_values in enumerate(rows):
        cells = table.add_row().cells
        for column_index in range(columns):
            value = row_values[column_index] if column_index < len(row_values) else ""
            cell = cells[column_index]
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, "1A5B93")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_run_with_inline_bold(
                paragraph,
                value,
                font_size=9.5,
                color=RGBColor(255, 255, 255) if row_index == 0 else TEXT,
            )
            if row_index == 0:
                for run in paragraph.runs:
                    run.font.bold = True
    doc.add_paragraph()


def add_quote(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.55)
    paragraph.paragraph_format.right_indent = Cm(0.25)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    border = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "1A5B93")
    border.append(left)
    paragraph._p.get_or_add_pPr().append(border)
    add_run_with_inline_bold(paragraph, text, font_size=11, color=MUTED)


def render_markdown_to_docx(source: Path, output: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document()
    style_document(doc)

    title_added = False
    expect_subtitle = False
    in_code = False
    code_lines: list[str] = []
    index = 0

    while index < len(lines):
        raw = lines[index]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                paragraph = doc.add_paragraph(style="Course Code")
                paragraph.paragraph_format.left_indent = Cm(0.45)
                for code_line in code_lines:
                    run = paragraph.add_run(code_line + "\n")
                    set_font(run, size=9, name="Consolas")
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if stripped.startswith("|"):
            table_rows, index = parse_table(lines, index)
            add_table(doc, table_rows)
            continue

        if stripped == "---":
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(4)
            border = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "D9E2F3")
            border.append(bottom)
            paragraph._p.get_or_add_pPr().append(border)
            index += 1
            continue

        if stripped.startswith("# "):
            if title_added:
                doc.add_page_break()
            title = doc.add_paragraph(style="Title")
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_run_with_inline_bold(title, stripped[2:], font_size=22, color=ACCENT)
            title_added = True
            expect_subtitle = True
            index += 1
            continue

        if expect_subtitle and stripped:
            subtitle = doc.add_paragraph()
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            subtitle_run = subtitle.add_run(stripped)
            set_font(subtitle_run, size=11, color=MUTED)
            expect_subtitle = False
            index += 1
            continue

        if stripped.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 1")
            add_run_with_inline_bold(paragraph, stripped[3:], font_size=16, color=ACCENT)
            index += 1
            continue

        if stripped.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            add_run_with_inline_bold(paragraph, stripped[4:], font_size=13, color=ACCENT)
            index += 1
            continue

        if stripped.startswith("#### "):
            paragraph = doc.add_paragraph(style="Heading 3")
            add_run_with_inline_bold(paragraph, stripped[5:], font_size=11.5, color=ACCENT)
            index += 1
            continue

        if stripped.startswith("> "):
            add_quote(doc, stripped[2:])
            index += 1
            continue

        bullet_match = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if bullet_match:
            add_bullet(doc, bullet_match.group(2), level=1 if len(bullet_match.group(1)) >= 2 else 0)
            index += 1
            continue

        number_match = re.match(r"^(\s*)(\d+[.)])\s+(.*)$", line)
        if number_match:
            add_numbered(
                doc,
                number_match.group(2),
                number_match.group(3),
                level=1 if len(number_match.group(1)) >= 2 else 0,
            )
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        paragraph = doc.add_paragraph()
        add_run_with_inline_bold(paragraph, stripped)
        index += 1

    if in_code and code_lines:
        paragraph = doc.add_paragraph(style="Course Code")
        for code_line in code_lines:
            run = paragraph.add_run(code_line + "\n")
            set_font(run, size=9, name="Consolas")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: render_xplore_course_docx.py INPUT.md OUTPUT.docx", file=sys.stderr)
        return 2
    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    render_markdown_to_docx(source, output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
